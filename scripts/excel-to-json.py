"""
Konverter Skovvangsvej Excel-data til JSON for intern.html.

Output:
  data/skovvangsvej.json — projektdata: KPI'er, kategorier, bilag
  data/baseline.json     — referenceværdier til beregneren

Køres ved at:
  python3 scripts/excel-to-json.py
"""

from openpyxl import load_workbook
from datetime import datetime
import json
import pathlib

ROOT = pathlib.Path(__file__).resolve().parents[1]
SRC = ROOT / "assets" / "project calculator" / "Skovvangsvej Færdigt projekt.xlsx"
OUT_PROJECT = ROOT / "data" / "skovvangsvej.json"
OUT_BASELINE = ROOT / "data" / "baseline.json"

# Kategori-typer (fra Estimering-arket — bruges af beregneren)
KATEGORI_TYPER = {
    "Materialer": "variabel",
    "Microcement": "variabel",
    "Gulve": "variabel",
    "Maler": "variabel",
    "Køkken": "fast",
    "Hårde Hvidevare": "fast",
    "Bad & Toilet": "fast",
    "Inventar": "fast",
    "VVS": "fast",
    "Forplejning": "fast",
    "Elektriker": "semi-variabel",
    "Diverse": "semi-variabel",
    "El": "semi-variabel",
}


def parse_dato(val):
    """Konvertér celle-værdi til ISO-dato (YYYY-MM-DD) hvis muligt."""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, str):
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(val, fmt).strftime("%Y-%m-%d")
            except ValueError:
                continue
        return val
    return None


def main():
    wb = load_workbook(SRC, data_only=True)

    # ---------- Pivot variable → kategori-totaler ----------
    pivot = wb["Pivot variable"]
    kategori_totaler = {}
    for row in pivot.iter_rows(min_row=4, values_only=True):
        navn, total = row[0], row[1]
        if navn and navn != "Grand Total" and total is not None:
            kategori_totaler[navn] = round(float(total), 2)

    # ---------- Claude variabel → bilag pr. kategori ----------
    bilag_ark = wb["Claude variabel"]
    bilag_pr_kategori = {k: [] for k in kategori_totaler}
    for row in bilag_ark.iter_rows(min_row=4, values_only=True):
        dato, leverandør, fakturanr, produkt, kategori, sek_kat, tilhører, antal, stk_pris, total_ex, moms, total_inkl = (
            row[:12] if len(row) >= 12 else (row + (None,) * (12 - len(row)))
        )
        if not kategori or kategori not in bilag_pr_kategori:
            continue
        if total_inkl is None:
            continue
        bilag_pr_kategori[kategori].append({
            "dato": parse_dato(dato),
            "leverandør": leverandør or "",
            "produkt": produkt or "",
            "antal": antal if antal else 1,
            "beløb": round(float(total_inkl), 2),
        })

    # Sortér bilag pr. kategori efter dato
    for k in bilag_pr_kategori:
        bilag_pr_kategori[k].sort(key=lambda b: b["dato"] or "")

    # ---------- Saml kategorier med type + bilag ----------
    kategorier = []
    for navn, total in sorted(kategori_totaler.items(), key=lambda x: -x[1]):
        kategorier.append({
            "navn": navn,
            "type": KATEGORI_TYPER.get(navn, "ukendt"),
            "total": total,
            "bilag": bilag_pr_kategori.get(navn, []),
        })

    # ---------- Lejlighedsinfo + Analyse → KPI ----------
    info = wb["Lejlighedsinfo"]
    analyse = wb["Analyse"]
    købspris = info["B5"].value
    salgspris = info["B6"].value
    projektvarighed = info["B9"].value
    tinglyst = info["E9"].value
    bbr = info["E10"].value
    ltv = analyse["B7"].value
    total_omkostninger = analyse["B16"].value
    fortjeneste_lav = analyse["E14"].value
    fortjeneste_est = analyse["G14"].value
    roi_lav = analyse["E16"].value
    roi_est = analyse["G16"].value
    renovering_total = kategori_totaler.get("Grand Total") or sum(kategori_totaler.values())

    projekt = {
        "id": "skovvangsvej",
        "navn": "Skovvangsvej 167",
        "adresse": "Skovvangsvej 167, st. tv., 8200 Aarhus N",
        "status": "Færdigt",
        "kpi": {
            "købspris": int(købspris),
            "salgspris": int(salgspris),
            "m2_tinglyst": int(tinglyst),
            "m2_bbr": int(bbr),
            "projektvarighed_mdr": int(projektvarighed),
            "antal_bad": 1,
            "stand_baseline": 4,
            "renovering_total": round(renovering_total, 2),
            "total_omkostninger": round(total_omkostninger, 2),
            "fortjeneste_estimeret": round(fortjeneste_est, 2),
            "fortjeneste_lav": round(fortjeneste_lav, 2),
            "roi_estimeret_pct": round(roi_est * 100, 2),
            "roi_lav_pct": round(roi_lav * 100, 2),
            "ltv": round(ltv, 4),
        },
        "kategorier": kategorier,
    }

    # ---------- Baseline til beregneren ----------
    # Aggregér pr. type baseret på Estimering-arket (ref m² = tinglyst, ref bad = 1)
    est = wb["Estimering"]
    faste, variable, semi = {}, {}, {}
    for r in range(6, 20):
        navn = est.cell(row=r, column=1).value
        beløb = est.cell(row=r, column=2).value
        type_ = est.cell(row=r, column=3).value
        if not navn or beløb is None or type_ is None:
            continue
        beløb = round(float(beløb), 2)
        if type_ == "Fast":
            faste[navn] = beløb
        elif type_ == "Variabel":
            variable[navn] = beløb
        elif type_ == "Semi-variabel":
            semi[navn] = beløb

    sum_fast = round(sum(faste.values()), 2)
    sum_var = round(sum(variable.values()), 2)
    sum_semi = round(sum(semi.values()), 2)
    ref_m2 = int(tinglyst)
    ref_bad = 1

    baseline = {
        "kilde": "Skovvangsvej (færdigt projekt)",
        "reference": {
            "m2": ref_m2,
            "antal_bad": ref_bad,
            "stand": 4,
            "scope_faktor": 1.0,
        },
        "totals": {
            "faste": sum_fast,
            "variable": sum_var,
            "semi_variable": sum_semi,
            "total": round(sum_fast + sum_var + sum_semi, 2),
        },
        "kategorier": {
            "faste": faste,
            "variable": variable,
            "semi_variable": semi,
        },
        "formler": {
            "faste":        "sum_faste * (antal_bad / ref_bad) * scope_faktor",
            "variable":     "sum_variable * (m2 / ref_m2)",
            "semi_variable": "sum_semi * (m2 / ref_m2) * scope_faktor",
            "stand_til_faktor": "stand≤4: 0.3 + (stand-1) * 0.7/3   ;   stand>4: 1.0 + (stand-4) * 1.0/6",
        },
        "risiko_grænser": {
            "ltv_max": 0.75,
            "renovering_pct_max": 0.15,
            "min_fortjeneste": 200000,
            "stress_prisfald_pct": 0.15,
        },
        "andre_omkostninger_pct": {
            "køb_pct_af_købspris": 0.02,
            "finansiering_pct_af_lånebeløb": 0.03,
            "salg_pct_af_salgspris": 0.025,
        },
    }

    OUT_PROJECT.parent.mkdir(parents=True, exist_ok=True)
    OUT_PROJECT.write_text(json.dumps(projekt, ensure_ascii=False, indent=2), encoding="utf-8")
    OUT_BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"✓ Skrev {OUT_PROJECT.relative_to(ROOT)}")
    print(f"   Kategorier: {len(kategorier)} | Total bilag: {sum(len(k['bilag']) for k in kategorier)}")
    print(f"   Total renovering: {renovering_total:,.0f} kr")
    print(f"✓ Skrev {OUT_BASELINE.relative_to(ROOT)}")
    print(f"   Faste: {sum_fast:,.0f}  Variable: {sum_var:,.0f}  Semi: {sum_semi:,.0f}  (ref m²={ref_m2})")


if __name__ == "__main__":
    main()
